"""Extração de texto dos formatos aceitos na base de conhecimento.

Formatos: TXT, Markdown, PDF, DOCX, CSV e XLSX. Qualquer falha de leitura é
convertida em :class:`UnprocessableFileError` com mensagem em português, sem
expor detalhes internos da biblioteca.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from app.core.errors import UnprocessableFileError

#: Limite de caracteres extraídos por documento (proteção de memória).
MAX_CHARACTERS = 400_000


@dataclass(slots=True)
class ExtractedText:
    """Texto extraído e metadados da extração."""

    text: str
    pages: int = 1
    truncated: bool = False
    notes: list[str] | None = None


def _truncate(text: str) -> tuple[str, bool]:
    if len(text) <= MAX_CHARACTERS:
        return text, False
    return text[:MAX_CHARACTERS], True


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnprocessableFileError("Não foi possível ler o arquivo: codificação não suportada.")


def extract_plain(content: bytes) -> ExtractedText:
    """TXT e Markdown."""
    text, truncated = _truncate(_decode(content))
    return ExtractedText(text=text, truncated=truncated)


def extract_pdf(content: bytes) -> ExtractedText:
    """PDF via pypdf (apenas texto; não faz OCR)."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependência do backend
        raise UnprocessableFileError("Suporte a PDF não disponível neste ambiente.") from exc

    try:
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise UnprocessableFileError(
            "Não foi possível ler o PDF. O arquivo pode estar protegido ou corrompido."
        ) from exc

    joined = "\n\n".join(
        f"[página {index}]\n{page.strip()}"
        for index, page in enumerate(pages, start=1)
        if page.strip()
    )
    if not joined.strip():
        raise UnprocessableFileError(
            "O PDF não contém texto selecionável. "
            "Documentos digitalizados exigem OCR, que não faz parte desta versão."
        )
    text, truncated = _truncate(joined)
    return ExtractedText(text=text, pages=len(pages), truncated=truncated)


def extract_docx(content: bytes) -> ExtractedText:
    """DOCX via python-docx (parágrafos e tabelas)."""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependência do backend
        raise UnprocessableFileError("Suporte a DOCX não disponível neste ambiente.") from exc

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise UnprocessableFileError("Não foi possível ler o arquivo DOCX.") from exc

    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    if not parts:
        raise UnprocessableFileError("O documento DOCX está vazio.")
    text, truncated = _truncate("\n\n".join(parts))
    return ExtractedText(text=text, truncated=truncated)


def extract_csv(content: bytes) -> ExtractedText:
    """CSV convertido em texto linha a linha (``coluna: valor``)."""
    raw = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.DictReader(io.StringIO(raw), delimiter=delimiter)
    lines: list[str] = []
    for index, row in enumerate(reader, start=1):
        pairs = [f"{key}: {value}" for key, value in row.items() if key and value]
        if pairs:
            lines.append(f"Registro {index} — " + "; ".join(pairs))
    if not lines:
        raise UnprocessableFileError("O CSV não possui linhas aproveitáveis.")
    text, truncated = _truncate("\n".join(lines))
    return ExtractedText(text=text, truncated=truncated)


def _label_for(labels: list[str], position: int) -> str:
    """Rótulo da coluna, ou um nome genérico quando o cabeçalho não existe."""
    if position < len(labels) and labels[position]:
        return labels[position]
    return f"coluna {position + 1}"


def extract_xlsx(content: bytes) -> ExtractedText:
    """XLSX convertido em texto por aba e linha."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependência do backend
        raise UnprocessableFileError("Suporte a XLSX não disponível neste ambiente.") from exc

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise UnprocessableFileError("Não foi possível abrir a planilha.") from exc

    blocks: list[str] = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        labels = [str(cell) if cell is not None else "" for cell in (header or ())]
        blocks.append(f"## Aba {sheet.title}")
        for index, row in enumerate(rows, start=2):
            pairs = [
                f"{_label_for(labels, position)}: {value}"
                for position, value in enumerate(row)
                if value not in (None, "")
            ]
            if pairs:
                blocks.append(f"Linha {index} — " + "; ".join(pairs))
    workbook.close()
    if len(blocks) <= 1:
        raise UnprocessableFileError("A planilha não possui dados aproveitáveis.")
    text, truncated = _truncate("\n".join(blocks))
    return ExtractedText(text=text, truncated=truncated)


EXTRACTORS = {
    ".txt": extract_plain,
    ".md": extract_plain,
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".csv": extract_csv,
    ".xlsx": extract_xlsx,
}


def extract_text(content: bytes, *, extension: str) -> ExtractedText:
    """Despacha a extração conforme a extensão do arquivo."""
    extractor = EXTRACTORS.get(extension.lower())
    if extractor is None:
        raise UnprocessableFileError(
            "Formato não suportado na base de conhecimento.",
            details={"extension": extension, "allowed": sorted(EXTRACTORS)},
        )
    return extractor(content)
